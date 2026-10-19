#!/usr/bin/env python3
"""Build a self-contained static web version of the FYP deck for display (Vercel).

Usage:  python3 docs/build-deck-site.py
Input:  docs/FYP-Presentation.pptx  (produced by docs/fyp-presentation.py)
Output: docs/deck/index.html        (single file, images inlined as data URIs)
        docs/deck/FYP-Presentation.pptx  (download copy)

The HTML is generated *from the pptx*, so the web deck always matches the
PowerPoint deck exactly — regenerate after any change to the generator.
"""

from __future__ import annotations

import base64
import html
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "docs" / "FYP-Presentation.pptx"
OUT_DIR = ROOT / "docs" / "deck"

EMU_PER_PX = 9525.0  # 96 px per inch
FONT = "Calibri, 'Segoe UI', system-ui, -apple-system, Arial, sans-serif"


def px(emu) -> float:
    return round(emu / EMU_PER_PX, 2)


def pt_px(pt) -> float:
    return round(pt * 96 / 72, 2)


def color_of(font, default=None):
    try:
        rgb = font.color.rgb
        return f"#{rgb}" if rgb is not None else default
    except Exception:
        return default


def fill_of(shape, default=None):
    try:
        if shape.fill.type is not None and shape.fill.fore_color.rgb is not None:
            return f"#{shape.fill.fore_color.rgb}"
    except Exception:
        pass
    return default


def line_of(shape):
    try:
        rgb = shape.line.color.rgb
        if rgb is None:
            return None
        w = shape.line.width.pt if shape.line.width is not None else 0.75
        return f"{pt_px(w)}px solid #{rgb}"
    except Exception:
        return None


ALIGN = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right", PP_ALIGN.JUSTIFY: "justify"}


def render_tf(tf, box_w: float | None = None) -> str:
    """Render a text frame (textbox or cell) to inner HTML."""
    parts = []
    for p in tf.paragraphs:
        size = 18.0
        color = "#333F50"
        spans = []
        for r in p.runs:
            sz = r.font.size.pt if r.font.size is not None else None
            if sz:
                size = sz
            c = color_of(r.font.color, None)
            if c:
                color = c
            txt = html.escape(r.text or "")
            if not txt:
                continue
            bold = " font-weight:bold;" if r.font.bold else ""
            spans.append(f'<span style="font-size:{pt_px(sz or size)}px;color:{c or color};{bold}">{txt}</span>')
        psize = size
        if not spans:
            spans = ["&#8203;"]  # keep empty spacer paragraphs (deck uses them for rhythm)
        align = ALIGN.get(p.alignment, "left")
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.15
        mb = pt_px(p.space_after.pt) if p.space_after is not None else 0
        parts.append(
            f'<div style="font-size:{pt_px(psize)}px;line-height:{ls};text-align:{align};'
            f'margin-bottom:{mb}px;min-height:{pt_px(psize) * ls:.1f}px">{"".join(spans)}</div>'
        )
    pad_t, pad_b = 0.0, 0.0
    try:
        pad_t = px(tf.margin_top or 0)
        pad_b = px(tf.margin_bottom or 0)
        pad_l = px(tf.margin_left or 0)
        pad_r = px(tf.margin_right or 0)
    except Exception:
        pad_l = pad_r = 9.6
    valign = ""
    if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
        valign = "display:flex;flex-direction:column;justify-content:center;"
    return (f'<div style="{valign}padding:{pad_t}px {pad_r}px {pad_b}px {pad_l}px;">'
            + "".join(parts) + "</div>")


def render_table(shape) -> str:
    tbl = shape.table
    cols = [px(c.width) for c in tbl.columns]
    rows_html = []
    for r, row in enumerate(tbl.rows):
        cells = []
        for c, cell in enumerate(row.cells):
            bg = fill_of(cell, "#FFFFFF")
            border = line_of(cell) or "1px solid #D7E0EA"
            cells.append(
                f'<td style="width:{cols[c]}px;background:{bg};border:{border};'
                f'vertical-align:middle;">{render_tf(cell.text_frame)}</td>'
            )
        rows_html.append(f"<tr>{''.join(cells)}</tr>")
    return (f'<table style="position:absolute;left:{px(shape.left)}px;top:{px(shape.top)}px;'
            f'width:{px(shape.width)}px;border-collapse:collapse;table-layout:fixed;">'
            f'<colgroup>{"".join(f"<col style=width:{w}px>" for w in cols)}</colgroup>'
            f"{''.join(rows_html)}</table>")


def render_slide(slide) -> str:
    out = []
    for shape in slide.shapes:
        x, y, w, h = px(shape.left), px(shape.top), px(shape.width), px(shape.height)
        box = f"position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;"
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img = shape.image
            b64 = base64.b64encode(img.blob).decode()
            out.append(f'<img src="data:{img.content_type};base64,{b64}" '
                       f'style="{box}object-fit:fill;" alt="">')
        elif shape.has_table:
            out.append(render_table(shape))
        elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            bg = fill_of(shape, "transparent")
            border = line_of(shape)
            radius = 0.0
            try:
                adj = shape.adjustments[0]
                if adj:
                    radius = adj * min(w, h)
            except Exception:
                pass
            style = box + f"background:{bg};border-radius:{radius:.1f}px;"
            if border:
                style += f"border:{border};"
            body = ""
            if shape.has_text_frame and shape.text_frame.text.strip():
                body = render_tf(shape.text_frame)
            out.append(f'<div style="{style}">{body}</div>')
        elif shape.has_text_frame:
            tf = shape.text_frame
            color = "#333F50"
            for p in tf.paragraphs:
                for r in p.runs:
                    c = color_of(r.font.color, None)
                    if c:
                        color = c
                        break
            out.append(f'<div style="{box}color:{color};font-family:{FONT};">{render_tf(tf)}</div>')
    return "".join(out)


def main() -> None:
    prs = Presentation(PPTX)
    W, H = px(prs.slide_width), px(prs.slide_height)
    slides = [render_slide(s) for s in prs.slides]
    n = len(slides)

    slide_divs = "\n".join(
        f'<section class="slide{" active" if i == 0 else ""}" data-i="{i}">{body}</section>'
        for i, body in enumerate(slides)
    )

    page = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>TimetableMaster — Final Year Project defense</title>
<meta name="description" content="Final Year Project defense deck: TimetableMaster, an intelligent timetable platform built in Rust and React.">
<style>
  * {{ box-sizing: border-box; }}
  html, body {{ margin:0; height:100%; background:#0B1B2B; font-family:{FONT}; overflow:hidden; }}
  #stage {{ position:fixed; inset:0; display:flex; align-items:center; justify-content:center; }}
  #deck {{ width:{W}px; height:{H}px; position:relative; background:#fff; transform-origin:center center;
           box-shadow:0 18px 60px rgba(0,0,0,.55); border-radius:2px; }}
  .slide {{ position:absolute; inset:0; display:none; background:#fff; overflow:hidden; }}
  .slide.active {{ display:block; }}
  #hud {{ position:fixed; left:0; right:0; bottom:0; height:46px; display:flex; align-items:center;
          gap:14px; padding:0 16px; background:rgba(8,18,32,.92); color:#C9D9E8; font-size:13px;
          border-top:1px solid rgba(240,162,2,.35); backdrop-filter:blur(6px); }}
  #hud button {{ background:#12263F; color:#E8EEF5; border:1px solid #274B70; border-radius:6px;
                 padding:5px 12px; font-size:14px; cursor:pointer; }}
  #hud button:hover {{ background:#1B3A5E; border-color:#F0A202; }}
  #counter {{ min-width:64px; text-align:center; font-variant-numeric:tabular-nums; }}
  #bar {{ flex:1; height:3px; background:#1B3A5E; border-radius:2px; overflow:hidden; }}
  #fill {{ height:100%; width:0%; background:#F0A202; transition:width .18s ease; }}
  #hint {{ opacity:.65; }}
  #dl {{ color:#F0A202; text-decoration:none; font-weight:600; }}
  #dl:hover {{ text-decoration:underline; }}
  #brand {{ font-weight:700; color:#fff; letter-spacing:.02em; }}
  @media (max-width:640px) {{ #hint {{ display:none; }} }}
  @media print {{
    html, body {{ overflow:visible; background:#fff; }}
    #hud {{ display:none; }}
    #stage {{ position:static; display:block; }}
    #deck {{ transform:none !important; width:auto; height:auto; box-shadow:none; }}
    .slide {{ display:block !important; position:relative; width:{W}px; height:{H}px;
              page-break-after:always; zoom:.85; }}
    @page {{ size:landscape; margin:0; }}
  }}
</style>
</head>
<body>
<div id="stage"><div id="deck">
{slide_divs}
</div></div>
<div id="hud">
  <span id="brand">TimetableMaster</span>
  <button id="prev" title="Previous (←)">◀</button>
  <span id="counter">1 / {n}</span>
  <button id="next" title="Next (→)">▶</button>
  <div id="bar"><div id="fill"></div></div>
  <a id="dl" href="FYP-Presentation.pptx" download>Download .pptx</a>
  <span id="hint">← → space · F fullscreen · Ctrl+P prints to PDF</span>
</div>
<script>
(function () {{
  var slides = Array.prototype.slice.call(document.querySelectorAll('.slide'));
  var n = slides.length, i = 0;
  var counter = document.getElementById('counter');
  var fill = document.getElementById('fill');
  var deck = document.getElementById('deck');

  function fit() {{
    var s = Math.min(window.innerWidth / {W}, (window.innerHeight - 46) / {H}) * 0.97;
    deck.style.transform = 'scale(' + s + ')';
  }}

  function go(k) {{
    i = Math.max(0, Math.min(n - 1, k));
    slides.forEach(function (el, j) {{ el.classList.toggle('active', j === i); }});
    counter.textContent = (i + 1) + ' / ' + n;
    fill.style.width = ((i + 1) / n * 100) + '%';
    try {{ history.replaceState(null, '', '#' + (i + 1)); }} catch (e) {{}}
  }}

  document.getElementById('prev').onclick = function () {{ go(i - 1); }};
  document.getElementById('next').onclick = function () {{ go(i + 1); }};
  document.addEventListener('keydown', function (e) {{
    if (e.key === 'ArrowRight' || e.key === 'PageDown' || e.key === ' ') {{ e.preventDefault(); go(i + 1); }}
    else if (e.key === 'ArrowLeft' || e.key === 'PageUp') {{ go(i - 1); }}
    else if (e.key === 'Home') {{ go(0); }}
    else if (e.key === 'End') {{ go(n - 1); }}
    else if (e.key === 'f' || e.key === 'F') {{
      if (document.fullscreenElement) document.exitFullscreen();
      else document.documentElement.requestFullscreen && document.documentElement.requestFullscreen();
    }}
  }});
  var tx = null;
  document.addEventListener('touchstart', function (e) {{ tx = e.touches[0].clientX; }}, {{passive:true}});
  document.addEventListener('touchend', function (e) {{
    if (tx === null) return;
    var dx = e.changedTouches[0].clientX - tx;
    if (Math.abs(dx) > 48) go(i + (dx < 0 ? 1 : -1));
    tx = null;
  }}, {{passive:true}});
  window.addEventListener('resize', fit);
  window.addEventListener('hashchange', function () {{
    var mm = location.hash.match(/^#(\\d+)$/);
    if (mm) go(parseInt(mm[1], 10) - 1);
  }});
  var m = location.hash.match(/^#(\\d+)$/);
  if (m) i = Math.max(0, Math.min(n - 1, parseInt(m[1], 10) - 1));
  fit(); go(i);
}})();
</script>
</body>
</html>
"""

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "index.html").write_text(page, encoding="utf-8")
    shutil.copyfile(PPTX, OUT_DIR / PPTX.name)
    size = (OUT_DIR / "index.html").stat().st_size
    print(f"wrote {OUT_DIR / 'index.html'}  ({size // 1024} KB, {n} slides, {W}x{H})")
    print(f"copied {PPTX.name} alongside for download")


if __name__ == "__main__":
    main()
